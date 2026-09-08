import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from pptx import Presentation
from pptx.util import Emu

PPTX = "ASTROVA_SIH2026_FINAL_6_SLIDE_PRESENTATION_REVISED.pptx"
EMU_IN = 914400.0

prs = Presentation(PPTX)
SW = prs.slide_width / EMU_IN
SH = prs.slide_height / EMU_IN
print(f"SLIDE: {SW:.3f} x {SH:.3f} in, slides={len(prs.slides)}")

def rects_overlap(a, b):
    ax1, ay1, ax2, ay2 = a
    bx1, by1, bx2, by2 = b
    ox = min(ax2, bx2) - max(ax1, bx1)
    oy = min(ay2, by2) - max(ay1, by1)
    if ox > 0.02 and oy > 0.02:
        return (ox, oy)
    return None

all_findings = []

for idx, slide in enumerate(prs.slides, 1):
    print(f"\n{'='*70}\nSLIDE {idx}\n{'='*70}")
    boxes = []
    min_font = 999
    small_texts = []

    for shape in slide.shapes:
        try:
            x = shape.left / EMU_IN; y = shape.top / EMU_IN
            w = shape.width / EMU_IN; h = shape.height / EMU_IN
        except Exception:
            continue
        rect = (x, y, x + w, y + h)
        text = ""
        fsize = None
        if shape.has_text_frame:
            text = shape.text_frame.text.strip()
            for p in shape.text_frame.paragraphs:
                for r in p.runs:
                    if r.font.size:
                        pt = r.font.size.pt
                        fsize = pt if fsize is None else min(fsize, pt)
        st = str(shape.shape_type)
        name = f"{st.split('(')[0].strip()}|{text[:28]!r}"

        # out of bounds
        if rect[0] < -0.01 or rect[1] < -0.01 or rect[2] > SW + 0.01 or rect[3] > SH + 0.01:
            all_findings.append((idx, "OUT-OF-BOUNDS", f"{name} rect=({x:.2f},{y:.2f},{rect[2]:.2f},{rect[3]:.2f})"))

        if fsize is not None and text:
            min_font = min(min_font, fsize)
            if fsize < 9:
                small_texts.append((fsize, text[:40]))
        elif text:
            # Check paragraph font size if run font size is missing
            for p in shape.text_frame.paragraphs:
                if p.font.size:
                    pt = p.font.size.pt
                    min_font = min(min_font, pt)
                    if pt < 9:
                        small_texts.append((pt, text[:40]))
                    break

        boxes.append((name, rect, bool(text), text, fsize))

    print(f"-- min font on slide: {min_font}pt; texts <9pt: {len(small_texts)}")
    for f, t in small_texts[:8]:
        print(f"     {f}pt  {t!r}")

    seen = set()
    for i in range(len(boxes)):
        for j in range(len(boxes)):
            if i == j: continue
            n1, r1, t1, x1, _ = boxes[i]
            n2, r2, t2, x2, _ = boxes[j]
            if not (t1 and i < j): continue
            ov = rects_overlap(r1, r2)
            if ov and ov[0] > 0.15 and ov[1] > 0.15:
                key = (n1, n2)
                if key in seen: continue
                seen.add(key)
                print(f"     OVERLAP {ov[0]:.2f}x{ov[1]:.2f}in: {n1[:40]}  <->  {n2[:40]}")
                all_findings.append((idx, "OVERLAP", f"{n1[:40]} <-> {n2[:40]} ({ov[0]:.2f}x{ov[1]:.2f}in)"))

print(f"\n{'='*70}\nFINDINGS SUMMARY ({len(all_findings)})\n{'='*70}")
for f in all_findings:
    print(f"  S{f[0]} [{f[1]}] {f[2]}")
